# clinic/processor/office_extract.py
"""
Extraction des documents Word (.docx) et PowerPoint (.pptx) pour le Studio.

Gemini lit nativement les PDF et les images, pas les fichiers Office. On en tire :
- un texte structuré (titres, paragraphes, tableaux en Markdown, notes de diapositives) ;
- pour PowerPoint, les liaisons entre formes (« A → B ») : un logigramme dessiné avec
  des formes et des connecteurs n'existe que sous cette forme ;
- les images intégrées (captures, logigrammes collés), envoyées telles quelles.
"""
import io
import logging
from typing import Any, Dict, List

logger = logging.getLogger(__name__)

MAX_IMAGES = 12
MIN_IMAGE_BYTES = 6_000          # ignore icônes et puces graphiques
MAX_TEXT_CHARS = 150_000

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def office_type(filename: str, content_type: str = "") -> str | None:
    name, ctype = (filename or "").lower(), (content_type or "").lower()
    if name.endswith(".docx") or ctype == DOCX_MIME:
        return "docx"
    if name.endswith(".pptx") or ctype == PPTX_MIME:
        return "pptx"
    return None


def _md_table(rows: List[List[str]]) -> str:
    rows = [[(c or "").replace("\n", " ").replace("|", "/").strip() for c in r] for r in rows if any((c or "").strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _extract_docx(data: bytes) -> Dict[str, Any]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    out: List[str] = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith(("heading", "titre")):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                out.append("#" * min(int(level) + 1, 4) + " " + text)
            elif "list" in style.lower() or "liste" in style.lower():
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            table = Table(child, doc)
            md = _md_table([[cell.text for cell in row.cells] for row in table.rows])
            if md:
                out.append(md)
    images = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                blob = rel.target_part.blob
                if len(blob) >= MIN_IMAGE_BYTES:
                    images.append(blob)
            except Exception:
                continue
    return {"text": "\n\n".join(out), "images": images[:MAX_IMAGES], "connections": 0}


def _extract_pptx(data: bytes) -> Dict[str, Any]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(data))
    out: List[str] = []
    images: List[bytes] = []
    total_links = 0
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh

    for idx, slide in enumerate(prs.slides, 1):
        shapes = list(walk(slide.shapes))
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        out.append(f"## Diapositive {idx}" + (f" : {title}" if title else ""))
        text_by_id: Dict[int, str] = {}
        ordered = sorted(shapes, key=lambda s: ((s.top or 0) // 200_000, s.left or 0))  # lecture haut → bas, gauche → droite
        for sh in ordered:
            if sh.has_text_frame and sh.text_frame.text.strip():
                text = " ".join(sh.text_frame.text.split())
                text_by_id[sh.shape_id] = text
                if text != title:
                    out.append("- " + text)
            if getattr(sh, "has_table", False) and sh.has_table:
                md = _md_table([[c.text for c in r.cells] for r in sh.table.rows])
                if md:
                    out.append(md)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = sh.image.blob
                    if len(blob) >= MIN_IMAGE_BYTES:
                        images.append(blob)
                except Exception:
                    pass
        # Connecteurs : début/fin rattachés à des formes → enchaînement du logigramme
        links = []
        for sh in shapes:
            el = sh._element
            if el.tag.rsplit("}", 1)[-1] != "cxnSp":
                continue
            st, en = el.find(f".//{ns}stCxn"), el.find(f".//{ns}endCxn")
            if st is None or en is None:
                continue
            a, b = text_by_id.get(int(st.get("id", -1))), text_by_id.get(int(en.get("id", -1)))
            label = " ".join(sh.text_frame.text.split()) if sh.has_text_frame and sh.text_frame.text.strip() else ""
            if a and b:
                links.append(f"- {a} → {b}" + (f" [{label}]" if label else ""))
        if links:
            total_links += len(links)
            out.append("Enchaînements (flèches du schéma) :")
            out.extend(links)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append("Notes : " + " ".join(notes.split()))
    return {"text": "\n".join(out), "images": images[:MAX_IMAGES], "connections": total_links}


def extract_office(data: bytes, kind: str) -> Dict[str, Any]:
    """{text, images, connections, has_diagram} ; lève ValueError si le fichier est illisible."""
    try:
        res = _extract_docx(data) if kind == "docx" else _extract_pptx(data)
    except Exception as e:
        raise ValueError(f"Fichier {kind.upper()} illisible : {e}")
    res["text"] = res["text"][:MAX_TEXT_CHARS]
    res["has_diagram"] = bool(res["images"]) or res["connections"] > 0
    return res


def office_parts(f: Dict[str, Any]) -> list:
    """Parts Gemini d'un fichier Office déjà extrait (f['office'])."""
    office = f.get("office") or {}
    label = "présentation PowerPoint" if f.get("type") == "pptx" else "document Word"
    parts: list = [f"=== Contenu extrait du {label} « {f.get('filename')} » ===\n{office.get('text') or '(aucun texte)'}"]
    if office.get("images"):
        from PIL import Image
        parts.append(f"Images intégrées dans « {f.get('filename')} » ({len(office['images'])}) :")
        for blob in office["images"]:
            try:
                img = Image.open(io.BytesIO(blob))
                img.load()
                if max(img.size) > 1024:
                    img.thumbnail((1024, 1024), Image.Resampling.LANCZOS)
                parts.append(img)
            except Exception:
                continue
    return parts
