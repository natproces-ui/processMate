"""
specifications_router.py — Génération de SFD depuis les procédures ProcessMate.

Préfixe : /api/orchestration/specifications
Persistance : table Supabase `specifications`
"""

import io
import json
import re
import uuid
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Dict, Any

import httpx
import openpyxl
from fastapi import APIRouter, HTTPException, Form, File, UploadFile
from fastapi.responses import HTMLResponse, Response
from google.genai import types as genai_types
from pydantic import BaseModel

from database.supabase_client import get_supabase
from sfd.generation import generate_sfd, chat_with_sfd
from sfd.data_transformer import build_sfd_context
from sfd.html_renderer import render_html
from sfd.word_renderer import render_docx
from sfd.themes import list_themes
from sfd.schema import SFDDocument

logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/api/orchestration/specifications",
    tags=["Spécifications SFD"],
)


# ─── MODÈLES ──────────────────────────────────────────────────────────────────

class GenerateRequest(BaseModel):
    scope_type: str                         # 'theme' | 'category' | 'subcategory' | 'procedures'
    scope_id: Optional[str] = None         # UUID taxonomy (requis sauf pour scope_type='procedures')
    procedure_ids: Optional[List[str]] = None  # requis si scope_type='procedures'
    title: Optional[str] = ""
    style: Optional[str] = "corporate_blue"


class ChatRequest(BaseModel):
    message: str
    chat_history: Optional[List[Dict[str, str]]] = []


class StyleRequest(BaseModel):
    style: str


# ─── HELPERS ──────────────────────────────────────────────────────────────────

def _get_descendant_ids(all_nodes: List[Dict], root_id: str) -> List[str]:
    """Retourne récursivement tous les IDs descendants d'un nœud taxonomique."""
    children = [n["id"] for n in all_nodes if n.get("parent_id") == root_id]
    ids = [root_id]
    for cid in children:
        ids.extend(_get_descendant_ids(all_nodes, cid))
    return ids


def _resolve_procedures(db, scope_type: str, scope_id: Optional[str], procedure_ids: Optional[List[str]]) -> List[Dict]:
    """Résout le scope en liste de rows workflow avec les colonnes nécessaires."""
    select_cols = "id, title, taxonomy_id, workflow_json, enrichments_json, procedure_metadata_json"

    if scope_type == "procedures":
        if not procedure_ids:
            raise HTTPException(400, "procedure_ids requis pour scope_type='procedures'")
        res = db.table("workflows").select(select_cols).in_("id", procedure_ids).execute()
        return res.data or []

    if not scope_id:
        raise HTTPException(400, f"scope_id requis pour scope_type='{scope_type}'")

    # Récupérer tous les nœuds, puis descendre récursivement
    all_nodes = db.table("process_taxonomy").select("id, parent_id, name, level").execute().data or []
    descendant_ids = _get_descendant_ids(all_nodes, scope_id)

    res = db.table("workflows").select(select_cols).in_("taxonomy_id", descendant_ids).execute()
    return res.data or []


def _resolve_scope_name(db, scope_type: str, scope_id: Optional[str]) -> str:
    """Retourne le nom du nœud taxonomique sélectionné."""
    if scope_type == "procedures":
        return "Sélection de procédures"
    if not scope_id:
        return scope_type
    row = db.table("process_taxonomy").select("name").eq("id", scope_id).execute().data
    return row[0]["name"] if row else scope_id


def _spec_summary(spec: Dict) -> Dict:
    """Version allégée d'un row specification (sans sfd_json complet)."""
    return {k: v for k, v in spec.items() if k != "sfd_json"}


# ─── EXTRACTION DES SOURCES EXTERNES ─────────────────────────────────────────

_BINARY_MIME: Dict[str, str] = {
    ".pdf":  "application/pdf",
    ".jpg":  "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png":  "image/png",
    ".webp": "image/webp",
    ".gif":  "image/gif",
}

_TEXT_STRIP_RE = re.compile(r"<[^>]+>")


def _excel_to_text(content: bytes) -> str:
    """Extrait le texte d'un fichier Excel (toutes les feuilles)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    parts: List[str] = []
    for sheet in wb.worksheets:
        parts.append(f"=== Feuille : {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line  = " | ".join(cells).strip(" |")
            if line:
                parts.append(line)
    return "\n".join(parts)


async def _url_to_text(url: str) -> str:
    """Récupère une URL et extrait le texte brut (sans balises HTML)."""
    async with httpx.AsyncClient(timeout=15.0, follow_redirects=True) as client:
        resp = await client.get(url, headers={"User-Agent": "ProcessMate SFD/1.0"})
        resp.raise_for_status()
        raw  = _TEXT_STRIP_RE.sub(" ", resp.text)
        clean = re.sub(r"\s+", " ", raw).strip()
        return clean[:25_000]  # plafonné à 25k chars par URL


async def _process_sources(
    source_files: List[UploadFile],
    source_urls:  List[str],
) -> tuple[str, List[Any]]:
    """
    Retourne (additional_sources_text, image_parts).

    - Excel → texte injecté dans additional_sources
    - PDFs / images → parts Gemini multimodal
    - URLs → texte injecté dans additional_sources
    """
    text_blocks: List[str] = []
    image_parts: List[Any] = []

    for upload in source_files:
        content  = await upload.read()
        filename = upload.filename or ""
        ext      = Path(filename).suffix.lower()
        mime     = _BINARY_MIME.get(ext)

        if mime:
            # PDF ou image → part binaire pour Gemini multimodal
            image_parts.append(genai_types.Part.from_bytes(data=content, mime_type=mime))
            logger.info(f"[sources] Fichier binaire : {filename} ({mime}, {len(content)} bytes)")
        elif ext in (".xlsx", ".xls", ".xlsm"):
            try:
                text = _excel_to_text(content)
                text_blocks.append(f"--- Source Excel : {filename} ---\n{text}")
                logger.info(f"[sources] Excel extrait : {filename}")
            except Exception as exc:
                logger.warning(f"[sources] Impossible de lire Excel {filename} : {exc}")
        else:
            logger.warning(f"[sources] Format non supporté ignoré : {filename}")

    for url in source_urls:
        url = url.strip()
        if not url:
            continue
        try:
            text = await _url_to_text(url)
            text_blocks.append(f"--- Source URL : {url} ---\n{text}")
            logger.info(f"[sources] URL récupérée : {url} ({len(text)} chars)")
        except Exception as exc:
            logger.warning(f"[sources] Impossible de récupérer {url} : {exc}")

    return "\n\n".join(text_blocks), image_parts


# ─── ENDPOINTS ────────────────────────────────────────────────────────────────

@router.get("/themes")
async def get_themes():
    """Retourne les thèmes visuels disponibles."""
    return {"success": True, "themes": list_themes()}


@router.get("")
async def list_specifications(
    status: Optional[str] = None,
    scope_type: Optional[str] = None,
):
    """Liste les spécifications générées, sans le sfd_json complet."""
    db = get_supabase()
    query = db.table("specifications").select(
        "id, scope_type, scope_id, scope_name, title, style, status, version, procedure_ids, generated_by, created_at, updated_at"
    ).order("created_at", desc=True)

    if status:
        query = query.eq("status", status)
    if scope_type:
        query = query.eq("scope_type", scope_type)

    res = query.execute()
    return {"success": True, "specifications": res.data or []}


@router.post("/generate")
async def generate_specification(
    scope_type:        str           = Form(...),
    scope_id:          Optional[str] = Form(None),
    procedure_ids:     Optional[str] = Form(None),   # JSON list encodé
    title:             Optional[str] = Form(""),
    style:             Optional[str] = Form("corporate_blue"),
    user_instructions: Optional[str] = Form(""),
    source_urls:       Optional[str] = Form(""),     # JSON list encodé
    source_files:      List[UploadFile] = File(default=[]),
):
    """
    Génère un SFD depuis les procédures + sources externes optionnelles.

    Accepte multipart/form-data :
    - scope_type, scope_id, procedure_ids (JSON), title, style : périmètre
    - user_instructions : message libre d'orientation de la génération
    - source_urls       : JSON list d'URLs à intégrer
    - source_files      : PDFs, images (multimodal Gemini), Excel (texte extrait)
    """
    db = get_supabase()

    # Désérialiser les listes encodées en JSON
    proc_ids_list: Optional[List[str]] = json.loads(procedure_ids) if procedure_ids else None
    urls_list:     List[str]           = json.loads(source_urls)   if source_urls   else []

    # 1. Résoudre scope → procédures
    procedures = _resolve_procedures(db, scope_type, scope_id, proc_ids_list)
    if not procedures:
        raise HTTPException(404, "Aucune procédure trouvée pour ce scope")

    scope_name = _resolve_scope_name(db, scope_type, scope_id)
    title      = title or scope_name

    logger.info(
        f"[specs/generate] scope={scope_type} name={scope_name} procs={len(procedures)} "
        f"files={len(source_files)} urls={len(urls_list)} "
        f"instructions={'oui' if user_instructions else 'non'}"
    )

    # 2. Traiter les sources externes
    additional_sources, image_parts = await _process_sources(source_files, urls_list)

    # 3. Transformer procédures → contexte textuel
    context = build_sfd_context(
        procedures         = procedures,
        scope_type         = scope_type,
        scope_name         = scope_name,
        additional_sources = additional_sources,
    )

    # 4. Générer le SFD via Gemini
    try:
        sfd = await generate_sfd(
            context           = context,
            scope_type        = scope_type,
            scope_name        = scope_name,
            procedure_count   = len(procedures),
            title             = title,
            user_instructions = user_instructions or "",
            image_parts       = image_parts or None,
        )
    except Exception as exc:
        logger.exception(f"[specs/generate] Erreur Gemini : {exc}")
        raise HTTPException(500, f"Erreur de génération : {exc}")

    # 5. Persister en Supabase
    now = datetime.utcnow().isoformat()
    row = {
        "id":            str(uuid.uuid4()),
        "scope_type":    scope_type,
        "scope_id":      scope_id,
        "scope_name":    scope_name,
        "procedure_ids": [str(p["id"]) for p in procedures],
        "title":         title,
        "sfd_json":      sfd.model_dump(),
        "style":         style or "corporate_blue",
        "status":        "draft",
        "version":       1,
        "chat_history":  [],
        "created_at":    now,
        "updated_at":    now,
    }

    res = db.table("specifications").insert(row).execute()
    if not res.data:
        raise HTTPException(500, "Échec de la sauvegarde")

    return {
        "success":         True,
        "specification":   res.data[0],
        "procedure_count": len(procedures),
    }


@router.get("/{spec_id}")
async def get_specification(spec_id: str):
    """Retourne le détail complet d'une spécification (avec sfd_json)."""
    db = get_supabase()
    res = db.table("specifications").select("*").eq("id", spec_id).execute()
    if not res.data:
        raise HTTPException(404, "Spécification introuvable")
    return {"success": True, "specification": res.data[0]}


@router.get("/{spec_id}/preview", response_class=HTMLResponse)
async def preview_specification(spec_id: str):
    """Retourne le rendu HTML de la spécification."""
    db = get_supabase()
    res = db.table("specifications").select("sfd_json, style").eq("id", spec_id).execute()
    if not res.data:
        raise HTTPException(404, "Spécification introuvable")

    row   = res.data[0]
    style = row.get("style", "corporate_blue")

    try:
        sfd  = SFDDocument(**row["sfd_json"])
        html = render_html(sfd, style=style)
        return HTMLResponse(content=html, status_code=200)
    except Exception as exc:
        logger.exception(f"[specs/preview] Erreur rendu HTML : {exc}")
        raise HTTPException(500, f"Erreur de rendu : {exc}")


@router.post("/{spec_id}/export")
async def export_specification(spec_id: str):
    """Génère et retourne le document Word (.docx) de la spécification."""
    db = get_supabase()
    res = db.table("specifications").select("sfd_json, style, title").eq("id", spec_id).execute()
    if not res.data:
        raise HTTPException(404, "Spécification introuvable")

    row   = res.data[0]
    style = row.get("style", "corporate_blue")
    title = row.get("title", "specification")

    try:
        sfd   = SFDDocument(**row["sfd_json"])
        docx  = render_docx(sfd, style=style)
        fname = title.replace(" ", "_").replace("/", "-")[:60]
        return Response(
            content     = docx,
            media_type  = "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers     = {"Content-Disposition": f'attachment; filename="{fname}.docx"'},
        )
    except Exception as exc:
        logger.exception(f"[specs/export] Erreur rendu Word : {exc}")
        raise HTTPException(500, f"Erreur d'export : {exc}")


@router.post("/{spec_id}/chat")
async def chat_specification(spec_id: str, body: ChatRequest):
    """Raffine la spécification via chat Gemini."""
    db = get_supabase()
    res = db.table("specifications").select("*").eq("id", spec_id).execute()
    if not res.data:
        raise HTTPException(404, "Spécification introuvable")

    spec = res.data[0]

    try:
        current_sfd = SFDDocument(**spec["sfd_json"])
        result = await chat_with_sfd(
            current_sfd  = current_sfd,
            user_message = body.message,
            chat_history = body.chat_history or spec.get("chat_history") or [],
            scope_name   = spec.get("scope_name", ""),
        )
    except Exception as exc:
        logger.exception(f"[specs/chat] Erreur Gemini : {exc}")
        raise HTTPException(500, f"Erreur de chat : {exc}")

    # Mettre à jour l'historique et le sfd_json
    history = (spec.get("chat_history") or []) + [
        {"role": "user",      "content": body.message},
        {"role": "assistant", "content": result["agent_message"]},
    ]

    now = datetime.utcnow().isoformat()
    db.table("specifications").update({
        "sfd_json":     result["sfd"].model_dump(),
        "chat_history": history,
        "updated_at":   now,
    }).eq("id", spec_id).execute()

    return {
        "success":          True,
        "agent_message":    result["agent_message"],
        "sections_modified": result["sections_modified"],
        "sfd_json":         result["sfd"].model_dump(),
    }


@router.patch("/{spec_id}/style")
async def update_style(spec_id: str, body: StyleRequest):
    """Change le thème visuel d'une spécification."""
    db = get_supabase()
    db.table("specifications").update({
        "style":      body.style,
        "updated_at": datetime.utcnow().isoformat(),
    }).eq("id", spec_id).execute()
    return {"success": True, "style": body.style}


@router.patch("/{spec_id}/status")
async def update_status(spec_id: str, status: str):
    """Change le statut (draft → validated → archived)."""
    valid = {"draft", "validated", "archived"}
    if status not in valid:
        raise HTTPException(400, f"Statut invalide. Valeurs : {valid}")
    db = get_supabase()
    db.table("specifications").update({
        "status":     status,
        "updated_at": datetime.utcnow().isoformat(),
    }).eq("id", spec_id).execute()
    return {"success": True, "status": status}


@router.delete("/{spec_id}")
async def delete_specification(spec_id: str):
    """Supprime définitivement une spécification."""
    db = get_supabase()
    db.table("specifications").delete().eq("id", spec_id).execute()
    return {"success": True}
